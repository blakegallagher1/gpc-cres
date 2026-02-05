"""
Document ingestion helpers for Deal Room uploads.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict

import pandas as pd
import pdfplumber
from docx import Document as DocxDocument


def _read_pdf_text(file_path: Path, max_pages: int = 5) -> str:
    with pdfplumber.open(str(file_path)) as pdf:
        pages = pdf.pages[:max_pages]
        return "\n".join(page.extract_text() or "" for page in pages).strip()


def _ocr_pdf_text(file_path: Path, max_pages: int = 5) -> str:
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        return ""
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            prefix = Path(temp_dir) / "page"
            subprocess.run(
                [
                    "pdftoppm",
                    "-f",
                    "1",
                    "-l",
                    str(max_pages),
                    "-r",
                    "300",
                    "-png",
                    str(file_path),
                    str(prefix),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            images = sorted(Path(temp_dir).glob("page-*.png"))
            if not images:
                images = sorted(Path(temp_dir).glob("page*.png"))
            chunks = []
            for image in images:
                result = subprocess.run(
                    ["tesseract", str(image), "stdout"],
                    check=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.DEVNULL,
                    text=True,
                )
                if result.stdout:
                    chunks.append(result.stdout.strip())
            return "\n".join(chunk for chunk in chunks if chunk).strip()
    except (OSError, subprocess.SubprocessError, ValueError):
        return ""


def _read_docx_text(file_path: Path) -> str:
    doc = DocxDocument(str(file_path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()


def _read_tabular_preview(file_path: Path) -> Dict[str, Any]:
    if file_path.suffix.lower() in {".csv", ".tsv"}:
        df = pd.read_csv(file_path)
    else:
        df = pd.read_excel(file_path)
    preview = df.head(10)
    return {
        "columns": list(preview.columns),
        "rows": preview.to_dict(orient="records"),
        "row_count": int(df.shape[0]),
    }


def extract_document(file_path: str, _mime_type: str | None = None) -> Dict[str, Any]:
    """
    Extracts text and structured previews from an uploaded document.
    Returns a dict suitable for ingestion_jobs.extracted_data.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    extracted: Dict[str, Any] = {"source_path": str(path), "text": "", "tables": []}

    if suffix == ".pdf":
        extracted["text"] = _read_pdf_text(path)
        if not extracted["text"]:
            extracted["text"] = _ocr_pdf_text(path)
    elif suffix in {".docx", ".doc"}:
        extracted["text"] = _read_docx_text(path)
    elif suffix in {".xlsx", ".xls", ".csv", ".tsv"}:
        extracted["tables"].append(_read_tabular_preview(path))
    else:
        extracted["text"] = f"Unsupported file type: {suffix}"

    extracted["classification"] = _classify_extracted_text(extracted.get("text", ""))
    extracted["underwriting_map"] = _auto_map_underwriting(extracted)
    return extracted


def _classify_extracted_text(text: str) -> Dict[str, Any]:
    if not text:
        return {"document_type": "unknown", "confidence": 0.0}
    lowered = text.lower()
    if "rent roll" in lowered:
        return {"document_type": "rent_roll", "confidence": 0.7}
    if "operating expenses" in lowered or "opex" in lowered:
        return {"document_type": "expenses", "confidence": 0.6}
    if "offering memorandum" in lowered or "offering memo" in lowered:
        return {"document_type": "offering_memo", "confidence": 0.6}
    return {"document_type": "general", "confidence": 0.4}


def _auto_map_underwriting(_extracted: Dict[str, Any]) -> Dict[str, Any]:
    """
    Minimal auto-mapping stub. Returns placeholders to be refined by agents.
    """
    return {
        "rent_growth": None,
        "exit_cap_rate": None,
        "noi": None,
        "notes": "Auto-mapping placeholder - refine with underwriting agent.",
    }
